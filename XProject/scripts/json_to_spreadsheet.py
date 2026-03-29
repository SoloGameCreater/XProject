#!/usr/bin/env python3
"""Convert JSON files into CSV or XLSX tables."""

from __future__ import annotations

import argparse
import csv
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DEFAULT_INPUT = Path("Assets/ExtraRes/Configs/DataJson/FishGame")
DEFAULT_OUTPUT = Path("Docs/Generated/FishGameTables")
SUPPORTED_FORMATS = ("csv", "xlsx")
HEADER_ALIASES: dict[str, dict[str, str]] = {
    "fishbait": {
        "BaitId": "鱼饵ID",
        "Name": "名称",
        "Type": "类型",
        "Quality": "品质",
        "BuyPriceCoin": "购买价格Coin",
        "ConsumePerCast": "每次抛竿消耗",
        "TargetTags": "目标鱼标签",
        "HookWeightBonus": "上钩权重加成",
        "HighLevelWeightBonus": "高等级鱼权重加成",
        "WaitTimeMultiplier": "等待时间倍率",
        "UnlockRodLevel": "解锁鱼竿等级",
        "EnabledPhase": "启用阶段",
        "CanPurchase": "可购买",
    },
    "fishdistancetier": {
        "TierId": "距离档位ID",
        "MinDistance": "最小距离",
        "MaxDistance": "最大距离",
        "CanBite": "可咬钩",
        "FishLevelWeightBonus": "鱼等级权重加成",
        "RarityWeightCommon": "普通稀有度权重",
        "RarityWeightUncommon": "少见稀有度权重",
        "RarityWeightRare": "稀有度权重",
        "RarityWeightLegendary": "传说稀有度权重",
        "WaitTimeMultiplier": "等待时间倍率",
    },
    "fishglobal": {
        "Id": "配置ID",
        "MinCastDistance": "最小抛竿距离",
        "NoBiteDistance": "无咬钩距离",
        "MaxCastDistance": "最大抛竿距离",
        "FightFailDistance": "战斗失败距离",
        "FightLineClampDistance": "战斗线长钳制距离",
        "AutoRetrieveSpeed": "自动收线速度",
        "WaitTimeMin": "最短等待时间",
        "WaitTimeMax": "最长等待时间",
        "EmptyHookCooldown": "空钩冷却时间",
        "CatchLineDistance": "捕获线长距离",
        "StarterRodId": "初始鱼竿ID",
        "StarterBaitId": "初始鱼饵ID",
        "StarterBaitCount": "初始鱼饵数量",
        "BaitPackCount": "鱼饵包数量",
        "OverLevelWeightScale": "超等级权重系数",
        "ReelDrainBonus": "收线消耗加成",
        "LockDrainBonus": "锁线消耗加成",
        "StruggleReleaseDrainScale": "挣扎放线消耗倍率",
        "ControlSuppressionDivisor": "控制压制换算除数",
        "SprintForceMul": "冲刺力度倍率",
        "RestForceMul": "休息力度倍率",
        "MaxCastChargeDuration": "最大抛竿蓄力时长",
        "CastingDelay": "抛竿延迟",
        "TensionRelaxDuration": "张力缓和时长",
        "ReelTensionSpeedMul": "收线张力速度倍率",
        "MinimumReelSpeed": "最低收线速度",
        "NormalCatchWindowOffset": "常规捕获窗口偏移",
        "FatigueCatchWindowOffset": "疲劳捕获窗口偏移",
        "BurstReadyStaminaRatio": "冲刺准备体力阈值",
        "BurstReadyDelayMin": "冲刺准备最短延迟",
        "BurstReadyDelayMax": "冲刺准备最长延迟",
        "BurstReelDangerWindow": "冲刺收线危险窗口",
        "BurstReelDangerTensionPressureMin": "冲刺收线危险张力下限",
        "BurstReelDangerTensionPressureMax": "冲刺收线危险张力上限",
        "BurstReelProgressCapRatio": "冲刺收线进度上限比例",
        "BurstReelResistancePenaltyMin": "冲刺收线阻力惩罚下限",
        "BurstReelResistancePenaltyMax": "冲刺收线阻力惩罚上限",
        "WeightStaminaLerpMin": "重量体力插值下限",
        "WeightStaminaLerpMax": "重量体力插值上限",
        "WeightResistanceLerpMin": "重量阻力插值下限",
        "WeightResistanceLerpMax": "重量阻力插值上限",
        "WeightEscapeLerpMin": "重量逃逸插值下限",
        "WeightEscapeLerpMax": "重量逃逸插值上限",
        "SteadyForceLerpWeight": "平稳阶段力度插值权重",
        "BurstTensionDrainScale": "冲刺阶段张力消耗倍率",
        "FatigueTensionLerpWeight": "疲劳阶段张力插值权重",
        "SteadyTensionDrainScale": "平稳阶段张力消耗倍率",
        "ReelResistanceBase": "收线阻力基础值",
        "ReelResistanceDrainScale": "收线阻力消耗倍率",
        "BurstReelStruggleScale": "冲刺收线挣扎倍率",
        "FatigueReelResistanceFloor": "疲劳收线阻力下限",
        "FatigueReelResistanceScale": "疲劳收线阻力倍率",
        "PullForceStaminaLerpMin": "拉力体力插值下限",
        "FatigueControlBonus": "疲劳控制加成",
        "RelaxBoostScale": "放松增益倍率",
        "EarlyReelSpeedMul": "前期收线速度倍率",
        "WaitTimeMultiplierFloor": "等待时间倍率下限",
        "ControlSuppressionMax": "控制压制上限",
        "SteadyDurationScale": "平稳阶段时长倍率",
        "BaitMismatchWeightScale": "鱼饵不匹配权重倍率",
    },
    "fishrod": {
        "RodId": "鱼竿ID",
        "Name": "名称",
        "BuyCostCoin": "购买价格Coin",
        "LineStrength": "线强度",
        "ReelSpeed": "收线速度",
        "ControlPower": "控制力",
        "SuppressPower": "压制力",
        "SupportedBaitTypeMask": "支持鱼饵类型掩码",
        "SupportedBaitQualityMax": "支持鱼饵最高品质",
        "RecommendFishLevelMax": "推荐鱼等级上限",
    },
    "fishspecies": {
        "FishId": "鱼种ID",
        "Species": "鱼种名称",
        "Level": "等级",
        "Rarity": "稀有度",
        "WeightMin": "最小重量",
        "WeightMax": "最大重量",
        "BaseStamina": "基础体力",
        "Resistance": "抵抗值",
        "EscapeSpeed": "逃逸速度",
        "BasePricePerKg": "每公斤基础售价",
        "PreferredBaitTags": "偏好鱼饵标签",
        "CatchWeight": "被钓中权重",
        "RecommendRodLevel": "推荐鱼竿等级",
        "StruggleDurationMin": "挣扎最短时长",
        "StruggleDurationMax": "挣扎最长时长",
        "RestDurationMin": "休息最短时长",
        "RestDurationMax": "休息最长时长",
        "SprintChance": "冲刺概率",
        "BaseDrainPerSec": "基础消耗每秒",
        "RecoveryPerSec": "恢复每秒",
    },
}


@dataclass
class TableData:
    headers: list[str]
    rows: list[dict[str, Any]]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert a JSON file or a directory of JSON files into CSV or XLSX tables."
    )
    parser.add_argument(
        "input_path",
        nargs="?",
        default=str(DEFAULT_INPUT),
        help=f"JSON file or directory. Default: {DEFAULT_INPUT}",
    )
    parser.add_argument(
        "-o",
        "--output-dir",
        default=str(DEFAULT_OUTPUT),
        help=f"Directory to place generated tables. Default: {DEFAULT_OUTPUT}",
    )
    parser.add_argument(
        "-f",
        "--format",
        choices=SUPPORTED_FORMATS,
        default="xlsx",
        help="Output format.",
    )
    parser.add_argument(
        "--indent-json-cells",
        type=int,
        default=None,
        help="Pretty-print nested JSON stored in a cell with the given indent.",
    )
    parser.add_argument(
        "--no-alias-row",
        action="store_true",
        help="Do not add a Chinese header alias row above the field names.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    input_path = Path(args.input_path)
    output_dir = Path(args.output_dir)

    json_files = collect_json_files(input_path)
    if not json_files:
        raise SystemExit(f"No JSON files found under: {input_path}")

    output_dir.mkdir(parents=True, exist_ok=True)

    tables = [(json_file, load_table(json_file, args.indent_json_cells)) for json_file in json_files]
    generated_files: list[Path]
    if args.format == "xlsx" and input_path.is_dir():
        workbook_name = input_path.name or "tables"
        destination = output_dir / f"{workbook_name}.xlsx"
        write_xlsx_workbook(destination, tables, include_alias_row=not args.no_alias_row)
        generated_files = [destination]
    else:
        generated_files = []
        for json_file, table in tables:
            destination = output_dir / f"{json_file.stem}.{args.format}"
            if args.format == "csv":
                write_csv(destination, table, alias_row=get_alias_row(json_file.stem, table.headers, not args.no_alias_row))
            else:
                write_xlsx(
                    destination,
                    table,
                    sheet_name=json_file.stem,
                    alias_row=get_alias_row(json_file.stem, table.headers, not args.no_alias_row),
                )
            generated_files.append(destination)

    print(f"input={input_path}")
    print(f"output_dir={output_dir}")
    print(f"format={args.format}")
    print(f"files={len(generated_files)}")
    for path in generated_files:
        print(path.as_posix())


def collect_json_files(input_path: Path) -> list[Path]:
    if input_path.is_file():
        return [input_path] if input_path.suffix.lower() == ".json" else []

    if not input_path.exists():
        raise SystemExit(f"Input path does not exist: {input_path}")

    return sorted(
        path
        for path in input_path.glob("*.json")
        if path.is_file() and not path.name.endswith(".json.meta")
    )


def load_table(json_file: Path, indent_json_cells: int | None) -> TableData:
    with json_file.open("r", encoding="utf-8") as stream:
        payload = json.load(stream)

    rows = normalize_rows(payload)
    flattened_rows = [flatten_row(row, indent_json_cells) for row in rows]

    headers: list[str] = []
    for row in flattened_rows:
        for key in row.keys():
            if key not in headers:
                headers.append(key)

    return TableData(headers=headers, rows=flattened_rows)


def normalize_rows(payload: Any) -> list[dict[str, Any]]:
    if isinstance(payload, list):
        if not payload:
            return []
        if all(isinstance(item, dict) for item in payload):
            return payload
        return [{"value": item} for item in payload]

    if isinstance(payload, dict):
        list_key = find_primary_list_key(payload)
        if list_key is None:
            return [payload]

        prefix = {key: value for key, value in payload.items() if key != list_key}
        rows: list[dict[str, Any]] = []
        for item in payload[list_key]:
            if isinstance(item, dict):
                row = dict(prefix)
                row.update(item)
            else:
                row = dict(prefix)
                row[list_key] = item
            rows.append(row)
        return rows

    return [{"value": payload}]


def find_primary_list_key(payload: dict[str, Any]) -> str | None:
    candidates = []
    for key, value in payload.items():
        if isinstance(value, list) and value:
            candidates.append((key, value))

    if not candidates:
        return None

    for key, value in candidates:
        if all(isinstance(item, dict) for item in value):
            return key

    return candidates[0][0]


def flatten_row(row: dict[str, Any], indent_json_cells: int | None) -> dict[str, Any]:
    flattened: dict[str, Any] = {}
    for key, value in row.items():
        _flatten_value(flattened, key, value, indent_json_cells)
    return flattened


def _flatten_value(
    target: dict[str, Any],
    key: str,
    value: Any,
    indent_json_cells: int | None,
) -> None:
    if isinstance(value, dict):
        if not value:
            target[key] = ""
            return
        for child_key, child_value in value.items():
            next_key = f"{key}.{child_key}" if key else child_key
            _flatten_value(target, next_key, child_value, indent_json_cells)
        return

    if isinstance(value, list):
        target[key] = json.dumps(value, ensure_ascii=False, indent=indent_json_cells)
        return

    target[key] = value


def write_csv(destination: Path, table: TableData, alias_row: list[str] | None = None) -> None:
    with destination.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=table.headers, extrasaction="ignore")
        if alias_row:
            writer.writerow(dict(zip(table.headers, alias_row)))
        writer.writeheader()
        for row in table.rows:
            writer.writerow({header: sanitize_cell(row.get(header)) for header in table.headers})


def write_xlsx(
    destination: Path,
    table: TableData,
    sheet_name: str,
    alias_row: list[str] | None = None,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = safe_sheet_name(sheet_name)
    populate_sheet(sheet, table, alias_row=alias_row)
    workbook.save(destination)


def write_xlsx_workbook(
    destination: Path,
    tables: list[tuple[Path, TableData]],
    include_alias_row: bool,
) -> None:
    workbook = Workbook()
    first_sheet = True

    for json_file, table in tables:
        if first_sheet:
            sheet = workbook.active
            first_sheet = False
        else:
            sheet = workbook.create_sheet()
        sheet.title = safe_sheet_name(json_file.stem)
        alias_row = get_alias_row(json_file.stem, table.headers, include_alias_row)
        populate_sheet(sheet, table, alias_row=alias_row)

    workbook.save(destination)


def populate_sheet(sheet, table: TableData, alias_row: list[str] | None = None) -> None:
    header_row_index = 2 if alias_row else 1
    data_start_row_index = header_row_index + 1
    sheet.freeze_panes = f"A{data_start_row_index}"
    sheet.auto_filter.ref = (
        f"A{header_row_index}:{get_column_letter(max(1, len(table.headers)))}"
        f"{max(header_row_index, len(table.rows) + header_row_index)}"
    )

    if alias_row:
        sheet.append(alias_row)

    if table.headers:
        sheet.append(table.headers)
        for row in table.rows:
            sheet.append([sanitize_cell(row.get(header)) for header in table.headers])
    else:
        sheet.append(["value"])

    apply_sheet_style(sheet, alias_row=alias_row)


def apply_sheet_style(sheet, alias_row: list[str] | None = None) -> None:
    alias_fill = PatternFill(fill_type="solid", fgColor="E2F0D9")
    header_fill = PatternFill(fill_type="solid", fgColor="D9E2F3")
    alias_font = Font(name="Arial", bold=True, color="385723")
    header_font = Font(name="Arial", bold=True)
    body_font = Font(name="Arial")

    if alias_row:
        for cell in sheet[1]:
            cell.font = alias_font
            cell.fill = alias_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    header_row_index = 2 if alias_row else 1
    for cell in sheet[header_row_index]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in sheet.iter_rows(min_row=header_row_index + 1):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for column_cells in sheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 48)


def sanitize_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return value


def safe_sheet_name(name: str) -> str:
    invalid_chars = set('[]:*?/\\')
    sanitized = "".join("_" if ch in invalid_chars else ch for ch in name).strip()
    if not sanitized:
        sanitized = "Sheet1"
    return sanitized[:31]


def get_alias_row(sheet_name: str, headers: list[str], include_alias_row: bool) -> list[str] | None:
    if not include_alias_row:
        return None

    alias_map = HEADER_ALIASES.get(sheet_name.lower())
    if not alias_map:
        return None

    return [alias_map.get(header, header) for header in headers]


if __name__ == "__main__":
    main()
